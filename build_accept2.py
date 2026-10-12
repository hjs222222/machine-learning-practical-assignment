# -*- coding: utf-8 -*-
"""
第二次阶段验收汇报（三、径流预测部分）—— 论文式图表规范
====================================================
- 图题：黑体、居中、不加粗，置于图下方
- 表题：宋体、居中、不加粗，置于表上方
- 每图配图释文字（真实数据）
数据源：output/results 10~15 号 CSV + output/figures fig11~fig20
"""
import os
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = r"D:\machine-learning-practical-assignment"
FIG = os.path.join(BASE, "output", "figures")
RES = os.path.join(BASE, "output", "results")
OUT = os.path.join(BASE, "机器学习课程实践_第二次阶段验收汇报.docx")

CONTENT_W = Cm(15.5)


# ---------------- 基础工具 ----------------
def _set_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def para(doc, text, size=10.5, bold=False, align=None, indent=True,
         before=0, after=6, line=1.4):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(21)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    _set_font(r, "宋体", size, bold)
    return p


def heading(doc, text, level=1):
    sizes = {1: 15, 2: 13, 3: 12}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    _set_font(r, "黑体", sizes.get(level, 12), True)
    return p


def fig_caption(doc, text):
    """图题：黑体、居中、不加粗、置于图下方"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    _set_font(r, "黑体", 10.5, False)
    return p


def tab_caption(doc, text):
    """表题：宋体、居中、不加粗、置于表上方"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    _set_font(r, "宋体", 10.5, False)
    return p


def caption_text(doc, text):
    """图释：正文小字"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(10)
    pf.line_spacing = 1.3
    r = p.add_run(text)
    _set_font(r, "宋体", 10, False)
    return p


def figure(doc, path, cap, explain=None, width=None):
    if not os.path.exists(path):
        para(doc, f"[缺失图片：{os.path.basename(path)}]", indent=False)
        return
    doc.add_picture(path, width=width or CONTENT_W)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_caption(doc, cap)
    if explain:
        caption_text(doc, explain)


def add_table(doc, df, cap=None, size=9, widths=None, col_names=None):
    cols = list(col_names) if col_names else list(df.columns)
    if cap:
        tab_caption(doc, cap)
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头行重复（跨页时表头延续）
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(tblHeader)
    hdr = t.rows[0].cells
    for i, cn in enumerate(cols):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(str(cn))
        _set_font(r, "黑体", size, True)
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row.tolist()):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            _set_font(r, "宋体", size, False)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for rr in t.rows:
            for c_, w_ in zip(rr.cells, widths):
                c_.width = w_
    return t


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def read_csv(name):
    p = os.path.join(RES, name)
    return pd.read_csv(p) if os.path.exists(p) else None


# ---------------- 文档 ----------------
def main():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "宋体"
    st.font.size = Pt(10.5)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    # ============ 封面 ============
    for _ in range(4):
        para(doc, "", indent=False)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run("《机器学习课程实践》"), "黑体", 22, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run("第二次阶段验收汇报"), "黑体", 26, True)
    for _ in range(3):
        para(doc, "", indent=False)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run("基于机器学习的径流预测（USGS 01047000 流域）"), "宋体", 14, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run("—— 径流多步预测（多步直接 + 多输出策略）"), "宋体", 12, True)
    for _ in range(3):
        para(doc, "", indent=False)
    for txt in ["汇报人：（姓名）", "学    号：（学号）", "验收日期：2026 年    月    日"]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run(txt), "宋体", 12, False)
    page_break(doc)

    # ============ 一、验收内容与完成情况 ============
    heading(doc, "一、验收内容与完成情况", 1)
    para(doc, "本次阶段验收内容为实验指导书第三部分「径流预测」：使用多步直接预测与多输出预测"
              "两种方式预测未来 1~7 天的径流，构建 2 种机器学习方法（人工神经网络、随机森林）、"
              "1 种深度学习方法（LSTM）与 2 种集成学习方法（简单平均、Stacking）的径流多步预测模型，"
              "利用 Grid Search CV 确定超参数并评价模型精度。")
    df = pd.DataFrame([
        ["多步直接预测", "对 h=1..7 每个预见期单独训练一个模型", "3 种方法 × 7 = 21 个"],
        ["多输出预测", "一个模型同时输出未来 1~7 天径流", "3 种方法 × 1 = 3 个"],
        ["个体模型合计", "ANN / 随机森林 / LSTM", "24 个"],
        ["集成学习", "简单平均、Stacking（两种策略各 8 个）", "16 个"],
        ["超参数寻优", "Grid Search CV + 时间序列交叉验证", "24 个模型逐个寻优"],
        ["精度评价", "RMSE/MAE/NSE/R²/PBIAS/MAPE，分预见期给出", "完成"],
    ], columns=["项目", "内容", "数量/状态"])
    add_table(doc, df, "表3-1  径流预测验收内容与完成情况", size=9.5)
    para(doc, "模型总数核算：个体模型 3 × 8 = 24 个；集成模型 2 × 8 = 16 个；合计 40 个模型，"
              "全部完成训练与精度评价。")
    page_break(doc)

    # ============ 二、方法调研 ============
    heading(doc, "二、适用于径流预测的方法调研", 1)
    para(doc, "径流预测是典型的时间序列回归问题，依据建模方法可分为三类：")
    para(doc, "（1）传统机器学习方法。人工神经网络（ANN/MLP）通过多层非线性映射逼近输入特征与"
              "径流之间的复杂关系，能拟合滞后、窗口特征中的非线性效应；随机森林（RF）是基于决策树"
              "的集成方法，对异常值稳健、可处理高维特征，且天然支持多输出回归。")
    para(doc, "（2）深度学习方法。LSTM 通过门控机制显式建模序列的长短期依赖，适合捕捉径流的"
              "记忆效应（本数据径流自相关 lag1=0.77、lag7=0.40，存在明显记忆）；输入为过去 30 天的"
              "特征序列，由网络自行学习时序模式。")
    para(doc, "（3）集成学习。简单平均直接对多个个体模型的预测取均值，降低方差、稳健性高；"
              "Stacking 以个体模型的预测为元特征，用 Ridge 回归学习最优融合权重，可进一步提升精度。"
              "调研表明，上述方法在洪水预报、径流预测文献中均有广泛应用，方法组合完整覆盖"
              "「2 种机器学习 + 1 种深度学习 + 2 种集成学习」的任务要求。")
    page_break(doc)

    # ============ 三、预测策略与特征选择 ============
    heading(doc, "三、多步预测策略与按策略特征选择", 1)

    heading(doc, "3.1 两种预测策略", 2)
    para(doc, "多步直接预测：对 h = 1,2,…,7 的每个预见期单独训练一个模型，第 h 个模型直接建立"
              "「当前特征 → 第 t+h 天径流」的映射 y(t+h) = f_h(X(t))，共 7 个模型；各模型相互独立，"
              "误差不向后传播。")
    para(doc, "多输出预测：只训练一个模型，输出层同时给出未来 7 天的径流 "
              "[y(t+1),…,y(t+7)] = F(X(t))，共 1 个模型；7 个输出共享隐层表示，可利用各预见期间的相关性。")
    para(doc, "数据划分严格按时间顺序：前 4 年（2000-01-01 ~ 2003-12-31）训练，"
              "第 5 年（2004-01-01 ~ 2004-12-31）测试，不打乱时序。")

    heading(doc, "3.2 按预测策略的特征选择（补充讲义要求）", 2)
    para(doc, "多步直接预测策略下，不同预见期对应的目标不同，最优特征集也应不同。"
              "本实验对每个预见期 h，在训练集上以第 t+h 天径流（对数）为目标用互信息法重新选择"
              "TOP50 特征，使各预见期模型使用各自的最优输入。")
    para(doc, "多输出预测策略下，采用补充讲义方法（3）：简单起见，选取第 t+1 天径流对应的特征"
              "作为全部输出的共用特征。")
    fs = read_csv("15_按预见期特征选择.csv")
    if fs is not None:
        add_table(doc, fs, "表3-2  按预见期选择的特征（TOP50 互信息）", size=9,
                  col_names=["预见期", "特征数", "Top5特征"])
    para(doc, "可以看到：预见期短时以径流自身的短期滞后与 3 天窗口特征为主（如 Discharge_lag1、"
              "Discharge_roll3_mean）；预见期延长后，长窗口特征与降水累积特征（Prcp_roll30_max）"
              "权重上升——这符合「预见期越长，越依赖前期水量累积信息」的水文规律。")
    figure(doc, os.path.join(FIG, "fig_流程.png"), "图3-1  径流多步预测总体流程",
           "总体流程为：数据分析 → 特征工程 → 按策略特征选择 → 个体模型（Grid Search CV 寻优）"
           "→ 集成融合 → 精度评价。")
    page_break(doc)

    # ============ 四、个体模型与超参数寻优 ============
    heading(doc, "四、个体模型构建与超参数寻优", 1)
    para(doc, "三种个体模型（ANN、随机森林、LSTM）均在多步直接与多输出两种策略下构建，"
              "并使用 Grid Search CV + 时间序列交叉验证（TimeSeriesSplit，3 折，前向链式切分）"
              "逐个模型确定最优超参数，避免普通 K 折造成时间泄漏。目标变换：径流取对数后建模，"
              "评价在原始尺度上进行。")
    bp = read_csv("12_最优超参数.csv")
    if bp is not None:
        bp2 = bp[["模型", "策略", "预见期", "最优超参数"]].copy()
        bp2["最优超参数"] = bp2["最优超参数"].str[:60]
        add_table(doc, bp2, "表3-3  各模型 Grid Search CV 最优超参数（摘要）", size=7,
                  widths=[Cm(4.0), Cm(2.2), Cm(1.4), Cm(7.6)])
    para(doc, "ANN 的最优结构为保守小结构（(64,)、(128,)、(64,32)），alpha 与学习率因预见期而异，"
              "避免大结构在样本量偏少时发生网格漂移过拟合；"
              "RF 的最优深度多为 10~15、叶节点最少 1~5 个样本；LSTM 最优为 64 或 64+32 单元、"
              "dropout=0.2/0.3、学习率 0.001/0.0005、训练轮数 80~120 轮（含早停）。")
    page_break(doc)

    # ============ 五、集成学习 ============
    heading(doc, "五、集成学习模型", 1)
    para(doc, "在两种预测策略下分别构建简单平均与 Stacking 两种集成模型，共 16 个。"
              "简单平均对三个个体模型的预测取算术平均；Stacking 以个体模型预测为元特征、"
              "Ridge 回归为元学习器，且元学习器的训练样本使用个体模型在训练期内的折外预测（OOF），"
              "防止融合权重被训练集过拟合污染。")
    sc = read_csv("14_Stacking元学习器系数.csv")
    if sc is not None:
        sc2 = sc.copy()
        sc2["Ridge系数"] = sc2["Ridge系数"].astype(str).str[:60]
        add_table(doc, sc2, "表3-4  Stacking 元学习器（Ridge）系数", size=8,
                  col_names=["策略", "预见期", "Ridge系数", "截距"])
    page_break(doc)

    # ============ 六、精度评价 ============
    heading(doc, "六、预测精度评价", 1)
    para(doc, "所有模型在完全相同的 359 个测试起始日上评价（2004-01-01 ~ 2004-12-24），"
              "保证各预见期样本数一致、结果可比。主要指标为水文预报通用的 NSE"
              "（Nash-Sutcliffe 效率系数，越接近 1 越好）、RMSE、R²、PBIAS 等。")

    heading(doc, "6.1 水文过程线（实测 vs 预测）", 2)
    figure(doc, os.path.join(FIG, "fig11_过程线_h1.png"), "图3-2  多步直接预测：预见期 1 天过程线",
           "各模型对 2004 年逐日径流的 1 天预见期预测与实测对比。")
    figure(doc, os.path.join(FIG, "fig11_过程线_h3.png"), "图3-3  多步直接预测：预见期 3 天过程线",
           "预见期 3 天时各模型预测相对实测的贴合程度。")
    figure(doc, os.path.join(FIG, "fig11_过程线_h7.png"), "图3-4  多步直接预测：预见期 7 天过程线",
           "预见期 7 天时预测平滑化明显，洪峰幅度被低估，符合长预见期信息衰减规律。")
    figure(doc, os.path.join(FIG, "fig12_多输出过程线_h1.png"), "图3-5  多输出预测：预见期 1 天过程线",
           "多输出策略下单一模型对 t+1 天的预测与实测对比。")
    figure(doc, os.path.join(FIG, "fig12_多输出过程线_h7.png"), "图3-6  多输出预测：预见期 7 天过程线",
           "多输出策略下同一模型对 t+7 天的预测与实测对比。")
    page_break(doc)

    heading(doc, "6.2 观测—预测散点图", 2)
    figure(doc, os.path.join(FIG, "fig13_散点_h1.png"), "图3-7  多步直接预测：预见期 1 天散点图",
           "散点越贴近 1:1 线表示预测越准确。")
    figure(doc, os.path.join(FIG, "fig13_散点_h3.png"), "图3-8  多步直接预测：预见期 3 天散点图", None)
    figure(doc, os.path.join(FIG, "fig13_散点_h7.png"), "图3-9  多步直接预测：预见期 7 天散点图", None)

    heading(doc, "6.3 指标随预见期的变化", 2)
    figure(doc, os.path.join(FIG, "fig14_NSE随预见期变化_直接.png"),
           "图3-10  多步直接预测：NSE 随预见期的变化",
           "所有模型 NSE 随预见期延长单调下降，1 天预见期最高、7 天最低。")
    figure(doc, os.path.join(FIG, "fig15_NSE随预见期变化_多输出.png"),
           "图3-11  多输出预测：NSE 随预见期的变化", None)
    figure(doc, os.path.join(FIG, "fig16_RMSE随预见期变化_直接.png"),
           "图3-12  多步直接预测：RMSE 随预见期的变化",
           "RMSE 随预见期延长而增大，反映长预见期不确定性上升。")
    page_break(doc)

    heading(doc, "6.4 模型对比与误差分析", 2)
    figure(doc, os.path.join(FIG, "fig17_NSE对比_直接.png"),
           "图3-13  多步直接预测：各模型 1~7 天汇总 NSE 对比", None)
    figure(doc, os.path.join(FIG, "fig17_NSE对比_多输出.png"),
           "图3-14  多输出预测：各模型 1~7 天汇总 NSE 对比", None)
    figure(doc, os.path.join(FIG, "fig19_洪水事件放大.png"),
           "图3-15  典型洪水事件放大分析（预见期 1 天）",
           "放大展示测试期最大两场洪水的预报细节：洪峰出现时机预测较准，洪峰量级存在低估。")
    figure(doc, os.path.join(FIG, "fig20_误差分析.png"),
           "图3-16  预测误差分析",
           "左图为各模型预测误差分布，右图为绝对误差随流量量级的变化：误差随流量增大而增大，"
           "说明洪峰阶段是预测的主要难点。")

    heading(doc, "6.5 汇总指标", 2)
    pooled = read_csv("11_模型汇总指标.csv")
    if pooled is not None:
        p = pooled[["模型", "策略", "预见期", "RMSE", "MAE", "NSE", "R2",
                    "PBIAS(%)", "MAPE(%)", "NSE_log", "NSE等级"]].copy()
        p = p[p["预见期"] == "1-7(汇总)"]
        add_table(doc, p, "表3-5  各模型 1~7 天汇总精度指标", size=8)
    para(doc, "随机森林与 Stacking 集成整体最优（多步直接 1~7 天汇总 NSE 约 0.20，多输出"
              "策略简单平均/Stacking 约 0.20~0.21）；1 天预见期 NSE 最高约 0.5（随机森林 0.526、"
              "Stacking 0.521、ANN 0.448）。所有模型在长预见期（h≥5）精度下降，"
              "与「未来降水未知、信息衰减」的物理规律一致。对数尺度指标（NSE_log，h=1 达 0.775）"
              "明显更高，说明模型对中低水过程拟合较好，主要误差来自少数特大洪水。")
    page_break(doc)

    # ============ 七、技术实现与总结 ============
    heading(doc, "七、技术实现与总结", 1)
    para(doc, "代码工程由数据、特征、模型、实验、可视化与报告生成等多个核心模块组成，"
              "通过 run_all.py 一键复现全部图表与数值结果（V6 定稿实验）：")
    for b in ["数据加载与字段释义、可视化分析（data_analysis.py）；",
              "特征构造、归一化与特征选择（feature_engineering.py，88 候选特征）；",
              "模型实现与超参数搜索（models.py，GridSearchCV + TimeSeriesSplit）；",
              "预测实验与集成融合（experiment_v6.py，含 V2~V5 迭代版本）；",
              "精度评价与结果可视化（evaluation.py / results_analysis.py）；",
              "实验报告与验收汇报自动生成（make_report.py / build_accept2.py）。"]:
        para(doc, "· " + b, indent=False)
    para(doc, "防数据泄漏措施：归一化统计量、特征选择评分、Stacking 元特征均只在训练集上估计；"
              "滞后与窗口特征只向过去取数；交叉验证使用时间序列切分。所有预测结果保存为 CSV/NPZ，"
              "全部图表为高分辨率 PNG，完全可追溯。")
    para(doc, "本次验收完成 24 个个体模型与 16 个集成模型的构建、超参数寻优与精度评价，"
              "全部达到验收要求。")

    doc.save(OUT)
    print("已生成:", OUT)


if __name__ == "__main__":
    main()
